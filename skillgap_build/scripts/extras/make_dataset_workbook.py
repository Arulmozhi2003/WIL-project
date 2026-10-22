"""Builds SkillGap_Final_Dataset.xlsx and SkillGap_Final_Dataset_job_ads.csv next to skillgap_build. Needs openpyxl."""
import json, sys
from pathlib import Path
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[2]  # skillgap_build
OUTDIR = ROOT.parent
sys.path.insert(0, str(ROOT))
from skillgap.core import load_documents, SKILL_TAXONOMY

docs = load_documents(ROOT / "data/documents")
raw = pd.read_csv(ROOT / "data/raw/seek_australia.csv", encoding="utf-8-sig").set_index("url")
questions = json.loads((ROOT / "data/evaluation/questions.json").read_text())
ROLES = ["Graduate Data Analyst", "Graduate Data Scientist", "Graduate Business/Data Analyst"]
SKILLS = sorted(SKILL_TAXONOMY)

rows = []
for d in sorted(docs, key=lambda d: (ROLES.index(d["role"]), d["doc_id"])):
    r = raw.loc[d["source_url"]]
    r = r.iloc[0] if isinstance(r, pd.DataFrame) else r
    val = lambda k: "" if pd.isna(r[k]) else str(r[k]).strip()
    rows.append({"doc_id": d["doc_id"], "role": d["role"], "job_title": d["title"], "company": val("company_name"),
                 "city": val("city"), "state_area": val("state"), "job_type": val("job_type"),
                 "salary_offered": val("salary_offered"), "seek_category": val("category"), "post_date": d["collected_date"], "experience_level": d["experience_level"],
                 "source_url": d["source_url"], "skills": "; ".join(d["skills"]), "n_skills": len(d["skills"]),
                 "word_count": len(d["text"].split()), "ad_text": d["text"]})
ads = pd.DataFrame(rows)
ads.to_csv(OUTDIR / "SkillGap_Final_Dataset_job_ads.csv", index=False, encoding="utf-8")

wb = Workbook()
F = "Arial"
H_FONT = Font(name=F, bold=True, color="FFFFFF"); H_FILL = PatternFill("solid", fgColor="1C5CAB")
BODY = Font(name=F, size=10); BOLD = Font(name=F, size=10, bold=True)

def header(ws, cols, widths):
    for i, (c, w) in enumerate(zip(cols, widths), 1):
        cell = ws.cell(row=1, column=i, value=c); cell.font = H_FONT; cell.fill = H_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"; ws.row_dimensions[1].height = 30

# --- README ---
ws = wb.active; ws.title = "README"
lines = [
    ("SkillGap final dataset", None),
    ("", None),
    ("What this is", "The 159 cleaned Australian job ads used as the SkillGap knowledge base, their skill tags, and the 54 labelled evaluation questions."),
    ("Source", "JobsPikr (2018). 30000 job postings from SEEK Australia [Data set]. Zenodo. https://doi.org/10.5281/zenodo.1220874"),
    ("Licence", "CC BY 4.0. Reuse requires attribution to the source above. Ad text was written by the advertisers; use for study and research."),
    ("Posting dates", "10-17 April 2018"),
    ("Raw source", "30,000 ads, 12 columns, 65.9 MB CSV (MD5 7d8136907b7f4d4aa9958ea676aa6247)"),
    ("Selection", "Role by job title -> senior titles removed (8,916) -> Business Analyst ads must name a data skill (142 removed) -> 181 -> cleaning (1 too short) -> 15 exact duplicates -> 6 reposts -> 159 ads"),
    ("Reproduce", "skillgap_build/scripts/build_seek_dataset.py and scripts/build_questions.py"),
    ("Skill tags", "Skills explicitly named in the ad, matched against a list of 47 skills and their spellings (skillgap/core.py). Automatic, not yet checked by hand: see data/annotation/skill_labels.csv."),
    ("", None),
    ("Sheets", None),
    ("Job_Ads", "One row per ad (159). ad_text is the cleaned text used by SkillGap: a header line (title, company, location, job type, salary) followed by the ad description."),
    ("Skill_Matrix", "One row per ad, one column per skill: 1 if the ad names the skill, else 0."),
    ("Skill_Profile", "Per role: number and share of ads naming each skill (formulas over Skill_Matrix)."),
    ("Questions", "54 evaluation questions: 33 dev (tuning) / 21 test (final score); 12 unanswerable. relevant_doc_ids are the correct ads."),
    ("", None),
    ("Job_Ads columns", None),
    ("doc_id", "Role prefix (DA / DS / BA) + SEEK job number"),
    ("role", "Target role family assigned from the job title"),
    ("job_title, company", "As advertised"),
    ("city, state_area, job_type, salary_offered, seek_category", "Original SEEK fields (salary is blank when not stated)"),
    ("post_date", "Date the ad was posted (YYYY-MM-DD)"),
    ("experience_level", "Graduate/entry wording · 0–1 year · Experience not stated · 2+ years. Numerical requirements override title wording; ranges use the upper bound. Unknown experience is not treated as zero experience."),
    ("source_url", "Original SEEK link (ads from 2018 are no longer live)"),
    ("skills, n_skills", "Skills named in the ad, separated by ';', and their count"),
    ("word_count", "Words in ad_text"),
    ("ad_text", "Cleaned full text: HTML removed, broken characters repaired"),
]
for i, (k, v) in enumerate(lines, 1):
    ws.cell(row=i, column=1, value=k).font = Font(name=F, size=14 if i == 1 else 10, bold=True)
    if v: c = ws.cell(row=i, column=2, value=v); c.font = BODY; c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.cell(row=i, column=1).alignment = Alignment(vertical="top", wrap_text=True)
ws.column_dimensions["A"].width = 34; ws.column_dimensions["B"].width = 110

# --- Job_Ads ---
ws = wb.create_sheet("Job_Ads")
cols = list(ads.columns)
header(ws, cols, [15, 28, 38, 30, 16, 26, 12, 24, 30, 11, 16, 38, 50, 9, 11, 90])
for r in ads.itertuples(index=False):
    ws.append(list(r))
for row in ws.iter_rows(min_row=2):
    for c in row: c.font = BODY; c.alignment = Alignment(vertical="top", wrap_text=(c.column in (13, 16)))
    row[15].alignment = Alignment(vertical="top", wrap_text=False)
ws.auto_filter.ref = ws.dimensions
for r in range(2, ws.max_row + 1): ws.cell(row=r, column=12).hyperlink = ws.cell(row=r, column=12).value

# --- Skill_Matrix ---
ws = wb.create_sheet("Skill_Matrix")
header(ws, ["doc_id", "role"] + SKILLS, [15, 28] + [7] * len(SKILLS))
for c in range(3, 3 + len(SKILLS)):
    ws.cell(row=1, column=c).alignment = Alignment(text_rotation=90, horizontal="center")
ws.row_dimensions[1].height = 120
for d in rows:
    have = set(d["skills"].split("; ")) if d["skills"] else set()
    ws.append([d["doc_id"], d["role"]] + [1 if s in have else 0 for s in SKILLS])
for row in ws.iter_rows(min_row=2):
    for c in row: c.font = BODY
    for c in row[2:]: c.alignment = Alignment(horizontal="center")
ws.auto_filter.ref = ws.dimensions
n = len(rows) + 1  # last data row in Skill_Matrix

# --- Skill_Profile (formulas) ---
ws = wb.create_sheet("Skill_Profile")
header(ws, ["skill"] + [x for role in ROLES for x in (f"{role.replace('Graduate ','')} ads naming it", "share")] + ["All ads naming it", "share"],
       [24, 14, 9, 14, 9, 14, 9, 14, 9])
ws.cell(row=2, column=1, value="Ads in role").font = BOLD
for j, role in enumerate(ROLES):
    c = ws.cell(row=2, column=2 + 2 * j, value=f'=COUNTIF(Skill_Matrix!$B$2:$B${n},"{role}")'); c.font = BOLD
ws.cell(row=2, column=8, value="=B2+D2+F2").font = BOLD
for i, skill in enumerate(SKILLS):
    r = 3 + i; col = get_column_letter(3 + i)
    ws.cell(row=r, column=1, value=skill).font = BODY
    for j, role in enumerate(ROLES):
        cc, sc = 2 + 2 * j, 3 + 2 * j
        ws.cell(row=r, column=cc, value=f'=SUMIFS(Skill_Matrix!${col}$2:${col}${n},Skill_Matrix!$B$2:$B${n},"{role}")').font = BODY
        s = ws.cell(row=r, column=sc, value=f"=IF({get_column_letter(cc)}$2=0,0,{get_column_letter(cc)}{r}/{get_column_letter(cc)}$2)")
        s.font = BODY; s.number_format = "0%"
    ws.cell(row=r, column=8, value=f"=SUM(Skill_Matrix!${col}$2:${col}${n})").font = BODY
    s = ws.cell(row=r, column=9, value=f"=IF(H$2=0,0,H{r}/H$2)"); s.font = BODY; s.number_format = "0%"
ws.auto_filter.ref = f"A2:I{2 + len(SKILLS)}"
note = ws.cell(row=4 + len(SKILLS), column=1, value="Counts are formulas over the Skill_Matrix sheet; sort or filter this table to rank skills per role.")
note.font = Font(name=F, size=9, italic=True)

# --- Questions ---
ws = wb.create_sheet("Questions")
header(ws, ["id", "split", "type", "role", "question", "n_relevant", "relevant_doc_ids", "label_rule"], [6, 7, 13, 28, 70, 10, 60, 45])
for q in questions:
    ws.append([q["id"], q["split"], q["type"], q["role"], q["question"], len(q["relevant_doc_ids"]), "; ".join(q["relevant_doc_ids"]), q["label_rule"]])
for row in ws.iter_rows(min_row=2):
    for c in row: c.font = BODY; c.alignment = Alignment(vertical="top", wrap_text=c.column in (5, 7))
ws.auto_filter.ref = ws.dimensions

wb.calculation.fullCalcOnLoad = True
out = OUTDIR / "SkillGap_Final_Dataset.xlsx"
wb.save(out)
print("saved", out, "| ads", len(ads), "| questions", len(questions), "| skills", len(SKILLS))
# expected values for a spot check of the formulas
prof = {role: {s: sum(1 for d in rows if d["role"] == role and s in d["skills"].split("; ")) for s in SKILLS} for role in ROLES}
print("expected SQL:", {r[9:]: prof[r]["SQL"] for r in ROLES}, "role sizes:", {r[9:]: sum(d['role'] == r for d in rows) for r in ROLES})
