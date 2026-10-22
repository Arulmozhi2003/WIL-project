"""Builds SkillGap_Preliminary_Overview.docx next to skillgap_build. Needs python-docx."""
import json, os, glob, statistics, collections, sys
from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = Path(__file__).resolve().parents[2]  # skillgap_build
OUT = ROOT.parent / "SkillGap_Preliminary_Overview.docx"
sys.path.insert(0, str(ROOT))
from skillgap.core import load_documents, role_skill_profile, SKILL_TAXONOMY

docs = load_documents(ROOT / "data/documents")
exp = json.loads((ROOT / "results/experiments.json").read_text())
questions = json.loads((ROOT / "data/evaluation/questions.json").read_text())
best = json.loads((ROOT / "results/best_config.json").read_text())
runs = {r["id"]: r for r in exp["runs"]}
words = [len(d["text"].split()) for d in docs]
raw_bytes = os.path.getsize(ROOT / "data/raw/seek_australia.csv")
doc_bytes = sum(os.path.getsize(f) for f in glob.glob(str(ROOT / "data/documents/*.json")))
per_role = collections.Counter(d["role"] for d in docs)
cities = collections.Counter(d["text"].split("(")[1].split(",")[0] for d in docs)
qtypes = collections.Counter((q["split"], q["type"]) for q in questions)

ACCENT = RGBColor(0x1C, 0x5C, 0xAB)
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
for side in ("left_margin", "right_margin"): setattr(sec, side, Cm(2.2))
sec.top_margin = sec.bottom_margin = Cm(2)

st = doc.styles
st["Normal"].font.name = "Calibri"; st["Normal"].font.size = Pt(10.5)
st["Normal"].element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
st["Normal"].paragraph_format.space_after = Pt(5)
for lvl, size in ((1, 15), (2, 12.5), (3, 11)):
    h = st[f"Heading {lvl}"]; h.font.name = "Calibri"; h.font.size = Pt(size); h.font.bold = True
    h.font.color.rgb = ACCENT if lvl < 3 else RGBColor(0x26, 0x26, 0x26)
    h.paragraph_format.space_before = Pt(14 if lvl == 1 else 10); h.paragraph_format.space_after = Pt(4)

def para(text="", bold_lead=None, italic=False, size=None, align=None):
    p = doc.add_paragraph()
    if bold_lead:
        r = p.add_run(bold_lead); r.bold = True
    r = p.add_run(text); r.italic = italic
    if size: r.font.size = Pt(size)
    if align: p.alignment = align
    return p

def bullets(items, style="List Bullet"):
    for it in items:
        p = doc.add_paragraph(style=style)
        if isinstance(it, tuple):
            r = p.add_run(it[0]); r.bold = True; p.add_run(it[1])
        else: p.add_run(it)
        p.paragraph_format.space_after = Pt(2)

def shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr(); s = OxmlElement("w:shd")
    s.set(qn("w:val"), "clear"); s.set(qn("w:color"), "auto"); s.set(qn("w:fill"), hex_fill); tcPr.append(s)

def table(header, rows, widths, caption=None, font=9.5):
    if caption:
        c = doc.add_paragraph(); r = c.add_run(caption); r.bold = True; r.font.size = Pt(9.5)
        c.paragraph_format.space_after = Pt(2); c.paragraph_format.keep_with_next = True
    t = doc.add_table(rows=1, cols=len(header)); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, h in enumerate(header):
        cell = t.rows[0].cells[i]; cell.text = ""; run = cell.paragraphs[0].add_run(h); run.bold = True
        run.font.size = Pt(font); shade(cell, "DCE8F7")
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""; run = cells[i].paragraphs[0].add_run(str(v)); run.font.size = Pt(font)
            if isinstance(v, str) and v.startswith("**"):
                run.text = v.strip("*"); run.bold = True
    for i, w in enumerate(widths): t.columns[i].width = Cm(w)
    for row in t.rows:
        for i, w in enumerate(widths): row.cells[i].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t

f3 = lambda v: "–" if v is None else f"{v:.3f}"

# ---------------- Title ----------------
p = doc.add_paragraph(); r = p.add_run("SkillGap"); r.bold = True; r.font.size = Pt(26); r.font.color.rgb = ACCENT
p.paragraph_format.space_after = Pt(0)
p = doc.add_paragraph(); r = p.add_run("How the prototype works: data, pipeline, app and first results"); r.font.size = Pt(14)
p = doc.add_paragraph(); r = p.add_run(f"WIL Project 84 · Preliminary team discussion document · 29 September 2026 · results run {exp['meta']['timestamp'][:10]}, dataset {exp['meta']['dataset_fingerprint']}")
r.font.size = Pt(9); r.italic = True; r.font.color.rgb = RGBColor(0x59, 0x59, 0x59)
p = doc.add_paragraph()
r = p.add_run("Internal working document for the team, not the submitted report. Parts of the code, data curation and this document were produced with AI assistance (Claude Code), which must be listed in the Condition 3 Bounded Use of AI Declaration.")
r.font.size = Pt(9); r.italic = True

# ---------------- 1 Summary ----------------
doc.add_heading("1. SkillGap in one page", 1)
para("SkillGap answers a graduate's questions about three roles (Graduate Data Analyst, Graduate Data Scientist, Graduate Business/Data Analyst) using real Australian job ads. It retrieves the ads that match the question, has a local AI model write an answer that cites those ads, refuses questions the ads cannot answer, and lists the skills employers ask for that the user does not yet have.")
table(["Item", "Value"], [
    ["Source dataset", "30,000 SEEK Australia job ads (JobsPikr 2018, Zenodo, CC BY 4.0)"],
    ["Raw file", f"seek_australia.csv: {len(open(ROOT/'data/raw/seek_australia.csv',encoding='utf-8-sig').readline().split(','))} columns, 30,000 rows, {raw_bytes/1e6:.1f} MB (0.066 GB)"],
    ["Knowledge base after cleaning", f"{len(docs)} ads · {sum(words):,} words · {doc_bytes/1e3:.0f} KB of JSON"],
    ["Ads per role", f"Data Analyst {per_role['Graduate Data Analyst']} · Business/Data Analyst {per_role['Graduate Business/Data Analyst']} · Data Scientist {per_role['Graduate Data Scientist']}"],
    ["Evaluation questions", f"{len(questions)} labelled (33 dev / 21 test; 12 unanswerable)"],
    ["Retriever chosen", "BM25 with stopwords, stemming, normalised score and a tuned refusal threshold"],
    ["Answer model", "Llama 3.2 3B via Ollama, running locally (2.0 GB)"],
    ["Headline result (test set)", f"Refuses {round(runs['BT']['test']['summary']['abstain_rate_unanswerable']*6)}/6 unanswerable questions (starter: 0/6); selection score {f3(runs['B0']['test']['summary']['selection_score'])} → {f3(runs['BT']['test']['summary']['selection_score'])}"],
    ["Code", "about 1,160 lines of Python in 12 files; 15 automated tests"],
], [5, 11.6], caption="Table 1. Key facts")

# ---------------- 2 Data ----------------
doc.add_heading("2. The data we used", 1)
doc.add_heading("2.1 Source", 2)
para("JobsPikr (2018). 30000 job postings from SEEK Australia [Data set]. Zenodo. https://doi.org/10.5281/zenodo.1220874. Licence CC BY 4.0: free to reuse with attribution. The ads were posted between 10 and 17 April 2018. Each row has 12 fields: category, city, company name, geo, job board, job description, job title, job type, post date, salary offered, state and URL. The file we downloaded matches Zenodo's published MD5 checksum, so anyone can confirm they have the identical file.")
para("Why this dataset: it is real Australian data, it has the full text of every ad, the original SEEK link and date, and a licence that allows reuse. We rejected a newer 2026 dataset (NextGig) because it had only about 5 relevant Australian ads, its descriptions were AI-written summaries, and it had no source links.", italic=False)
para("Size, in plain terms: the whole raw file is 66 MB (about 0.07 GB). We only use 159 ads from it, which is under half a megabyte of text, about 60,000 words or roughly 240 pages of a novel. No gigabyte-scale data is involved; the only gigabyte-sized download is the answer-writing AI model (Section 2.4).")

doc.add_heading("2.2 From 30,000 ads to 159: the selection funnel", 2)
table(["Step", "Rule", "Ads left"], [
    ["Start", "Every posting in the file", "30,000"],
    ["Role by job title", "Keep titles that match one of the three role families (e.g. 'data scientist', 'reporting analyst', 'business analyst'); 20,761 other titles removed", "–"],
    ["Graduate focus", "Remove senior titles (senior, lead, principal, manager, head, director, architect, consultant, specialist…): 8,916 removed", "–"],
    ["Business/Data Analyst check", "A business analyst ad must name at least one data skill (SQL, Excel, Power BI…): 142 removed", "181"],
    ["Cleaning", "Strip HTML, repair broken characters ('â€“' → '–'), drop ads under 300 characters (1)", "180"],
    ["Exact duplicates", "Identical text posted twice (15)", "165"],
    ["Reposts", "Same title and at least 80% identical wording, e.g. one ad per city (6)", "**159**"],
], [3.3, 11, 2.3], caption="Table 2. Data selection, reproduced by scripts/build_seek_dataset.py")

doc.add_heading("2.3 What the knowledge base looks like", 2)
bullets([
    ("Ad length: ", f"median {statistics.median(words):.0f} words (shortest {min(words)}, longest {max(words)})."),
    ("Locations: ", ", ".join(f"{c} {n}" for c, n in cities.most_common(6)) + ", plus others."),
    ("Skills: ", f"each ad is tagged with the skills it explicitly names, using a list of {len(SKILL_TAXONOMY)} skills and their spellings (e.g. 'PowerBI', 'Power-BI' and 'Microsoft Power BI' all count as Power BI). On average an ad names {sum(len(d['skills']) for d in docs)/len(docs):.1f} skills."),
    ("One ad = one JSON file ", "in data/documents/, holding the id, role, title, full text, skills, SEEK URL, source type and date."),
])
rows = []
for role in ["Graduate Data Analyst", "Graduate Data Scientist", "Graduate Business/Data Analyst"]:
    prof = role_skill_profile(docs, role)
    rows.append([role.replace("Graduate ", ""), per_role[role], ", ".join(f"{p['skill']} {p['share']:.0%}" for p in prof[:7])])
table(["Role", "Ads", "Most requested skills (share of that role's ads)"], rows, [3.6, 1.2, 11.8], caption="Table 3. Skill profile per role")

doc.add_heading("2.4 Models and other files", 2)
table(["Item", "Size", "Used for"], [
    ["Llama 3.2 3B (Ollama)", "2.0 GB", "Writing the answer from the retrieved ads"],
    ["all-MiniLM-L6-v2 (sentence-transformers)", "about 90 MB (175 MB cache)", "Dense 'meaning-based' retrieval, used only as a comparison"],
    ["questions.json", f"{len(questions)} questions", "Measuring retrieval quality"],
    ["label_audit.csv", "72 rows", "Human check of our question labels (to do)"],
    ["skill_labels.csv", "20 ads", "Human check of skill extraction (to do)"],
], [5.5, 3.6, 7.5], caption="Table 4. Models and supporting files")

# ---------------- 3 Pipeline ----------------
doc.add_heading("3. How the code works, step by step", 1)
para("The system has two stages. The offline stage prepares the data once. The online stage runs every time someone asks a question.")
doc.add_heading("Offline stage (run once)", 2)
steps_off = [
    ("Step 1. Build the dataset. ", "scripts/build_seek_dataset.py downloads the SEEK file, verifies the checksum, applies the selection rules in Table 2 and hands the selected rows to ingest.py."),
    ("Step 2. Clean and tag. ", "ingest.py cleans the text, removes duplicates and reposts, tags every ad with the skills it names (extract_skills in skillgap/core.py), and writes one JSON file per ad plus a provenance record (data/provenance/MANIFEST.csv) linking each file to its original row."),
    ("Step 3. Build the questions. ", "scripts/build_questions.py writes 54 test questions with the ads that correctly answer each one, and splits them into dev (for tuning) and test (for the final score)."),
    ("Step 4. Tune and evaluate. ", "run_experiments.py tries each retrieval setting on the dev questions, keeps the best, scores all of them once on test, and saves the winner to results/best_config.json, which the app reads."),
]
bullets(steps_off)
doc.add_heading("Online stage (every question)", 2)
steps_on = [
    ("Step 5. Filter by role. ", "Only ads for the selected role are considered, so a Data Scientist question never returns a Business Analyst ad."),
    ("Step 6. Turn the question into keywords. ", "The question is lower-cased and split into words; filler words such as 'the', 'is' and 'what' are dropped; words are reduced to a stem so 'requested', 'requests' and 'request' match."),
    ("Step 7. Score every ad with BM25. ", "BM25 is the standard search-engine formula. An ad scores higher when it contains the question's words, especially rare, specific ones (like 'Tableau') and when it uses them repeatedly, with long ads penalised so they do not win just by being long."),
    ("Step 8. Decide whether to answer at all. ", "The score is converted to a 0–1 value: how much of the question's important vocabulary the ad covers. If even the best ad scores below the tuned threshold (0.128), SkillGap refuses instead of guessing. This is what makes 'What will salaries be in 2030?' get no answer."),
    ("Step 9. Pick the evidence. ", "The top 3 ads above the threshold become the evidence."),
    ("Step 10. Write the answer. ", "The 3 ads, the role and the question are sent to Llama 3.2 running on the laptop (nothing leaves the machine). The instructions tell it to use only these ads, end every claim with a citation such as [SOURCE DA_35991436], ignore any instructions hidden in the ads, and reply with a fixed refusal sentence if the ads do not answer the question."),
    ("Step 11. Check the answer. ", "The app reads the citations and warns if the answer cites an ad that was not retrieved, or has no citations at all."),
    ("Step 12. Compute the skill gaps. ", "Separately from the AI, the app counts how many of the role's ads name each skill and removes the skills the user already has. This is plain counting, so the AI cannot invent a skill gap, and the list does not change with the wording of the question."),
]
bullets(steps_on)

# ---------------- 4 Streamlit ----------------
doc.add_heading("4. The Streamlit app and how it answers a question", 1)
para("Streamlit turns a Python script (app.py) into a web page. It came with the project's starter kit. Run 'streamlit run app.py' and open http://localhost:8501. The page has:")
bullets([
    ("Inputs: ", "target role, your existing skills (comma-separated), your question, and a Settings panel (retriever, number of ads, AI model, prompt version, minimum share for a skill gap)."),
    ("Answer: ", "the AI-written answer with citations, or a message that the knowledge base cannot answer the question."),
    ("Potential skill gaps: ", "a table of skills you did not list, with 'named in N of M ads' and the share."),
    ("Retrieved evidence: ", "the ads used, each expandable to show the full text, its SEEK link and date."),
])
doc.add_heading("Worked example", 3)
para("Role: Graduate Data Analyst · Skills: Python, Excel · Question: 'Is experience building dashboards in Tableau requested?'", italic=True)
bullets([
    "Keywords after steps 6–7: experienc, build, dashboard, tableau, request. BM25 ranks ads DA_35991436, DA_35994633 and DA_35968969 highest; all pass the threshold.",
    "Llama receives the three ads and returns, in about 12 seconds: \"According to [SOURCE DA_35991436], creating intuitive dashboards … is a key responsibility, but it does not explicitly state that experience building dashboards in Tableau is required. However, [SOURCE DA_35994633] explicitly states that the ability to help customers realise the benefits of Tableau … is a required attribute …\"",
    "Citation check: both cited ads were retrieved, so no warning.",
    "Skill gaps for this user include Communication (39 of 80 ads), SQL (38 of 80) and Tableau (21 of 80).",
])
para("Asking 'What will the average salary for this role be in 2030?' instead stops at step 8: no ad clears the threshold, so the AI is never called and the page says the question cannot be answered from the knowledge base.")

# ---------------- 5 Evaluation ----------------
doc.add_heading("5. How we measured it, and the results", 1)
doc.add_heading("5.1 The questions", 2)
para(f"54 questions, 18 per role, in three types. Skill questions ({qtypes[('dev','skill')]+qtypes[('test','skill')]}) ask whether a skill is requested, worded differently from the ads (e.g. 'Do employers expect me to write database queries in SQL?'); the correct ads are all ads of that role that name the skill. Specific-ad questions ({qtypes[('dev','ad')]+qtypes[('test','ad')]}) name a job and company; the correct ad is that one. Unanswerable questions ({qtypes[('dev','unanswerable')]+qtypes[('test','unanswerable')]}) ask about things no ad mentions, checked automatically. 60% are 'dev' questions used for tuning; 40% are 'test' questions scored once at the end, so the final numbers are not tuned to the questions they are reported on.")
doc.add_heading("5.2 The measures", 2)
bullets([
    ("Hit@3: ", "share of answerable questions where at least one correct ad is in the top 3."),
    ("MRR: ", "rewards putting the first correct ad high (1st = 1, 2nd = 0.5, 3rd = 0.33)."),
    ("nDCG@3: ", "overall ranking quality of the top 3, from 0 to 1."),
    ("Decision accuracy: ", "share of all questions where the system correctly chose to answer or to refuse."),
    ("Refusal rate (unanswerable) / false refusals: ", "how often it refuses when it should, and wrongly refuses when it should answer."),
    ("Selection score: ", "average of nDCG@3 and decision accuracy; used to pick the winning setting on dev."),
])
doc.add_heading("5.3 Results: one change at a time", 2)
para("Each row adds one change to the row above, so its effect can be seen on its own. The winner (BT) was chosen on dev only; the table shows the test scores.")
labels = {"B0": "Starter version (original code)", "B1": "+ drop filler words", "B2": "+ stemming", "B3": "+ refusal threshold (raw score)",
          "B4": "+ threshold on 0–1 score", "BT": "Fully tuned BM25 (chosen)", "D0": "Meaning-based (MiniLM), no threshold",
          "D1": "Meaning-based + threshold", "H1": "Hybrid of BT and D1"}
rows = []
for cid in ["B0", "B1", "B2", "B3", "B4", "BT", "D0", "D1", "H1"]:
    s = runs[cid]["test"]["summary"]
    lab = labels[cid]; lab = f"**{lab}" if cid == "BT" else lab
    rows.append([cid, lab, f3(s["hit_rate_at_k"]), f3(s["mean_ndcg_at_k"]), f3(s["decision_accuracy"]),
                 f"{round(s['abstain_rate_unanswerable']*6)}/6", f3(s["selection_score"])])
table(["ID", "Configuration", "Hit@3", "nDCG@3", "Decision acc.", "Refused (unans.)", "Selection"], rows,
      [1.1, 5.9, 1.6, 1.7, 2.2, 2.3, 1.8], caption="Table 5. Held-out test results (21 questions)", font=9)
ci = exp["final_test_ci95"]
para(f"95% confidence interval for the chosen system's selection score: {f3(ci['selection_score'][0])} to {f3(ci['selection_score'][1])}. The interval is wide because the test set is small (21 questions), which is itself a finding to report.")
fig = ROOT / "results/figures/threshold_sweep.png"
if fig.exists():
    doc.add_picture(str(fig), width=Cm(14)); doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    c = para("Figure 1. Effect of the refusal threshold on dev questions. Too low and the system answers everything; too high and it refuses good questions. The dashed line is the chosen value.", size=9, align=WD_ALIGN_PARAGRAPH.CENTER); c.runs[0].italic = True
doc.add_heading("5.4 What the results say", 2)
bullets([
    "The biggest gain came from teaching the system to refuse: the starter version answered every unanswerable question; the tuned one refuses half of them on test without refusing any answerable one.",
    "Stemming and dropping filler words improved ranking (nDCG 0.820 → 0.882 at B2).",
    "Plain keyword search (BM25) beat the meaning-based model (MiniLM) on this data: job-ad questions hinge on exact tool names such as SQL or Tableau, which keyword search handles well.",
    "Combining the two (hybrid) did not help, partly because it answers whenever either method is confident, which weakens refusals.",
    "On test, one intermediate setting (B4, 0.886) scored slightly above the chosen one (0.862). We keep BT because it was chosen before looking at test; switching now would be cheating the held-out test.",
])

# ---------------- 6 Failures ----------------
doc.add_heading("6. Problems we found and fixed", 1)
table(["Problem", "How we found it", "Fix"], [
    ["Starter answered every question, even 'salary in 2030'", "It matched on 'this' and 'job'", "Filler-word removal + tuned refusal threshold"],
    ["Skill gaps changed with the question's wording", "Code review", "Gaps counted over all ads of the role"],
    ["C++ / C# lost; 'Power-BI' not recognised", "Unit tests", "New tokeniser and skill-matching rules"],
    ["First threshold overfitted (failed a test question)", "Smoke run", "Threshold placed mid-way in the best range"],
    ["Broken characters ('â€“'), 15 duplicates, 6 reposts", "Inspecting the real data", "Encoding repair and duplicate detection"],
    ["Plain 'pytest' could not find the code", "Teammate run", "Added pytest.ini"],
    ["Ollama returned 404", "App run", "The model had not been downloaded; ran 'ollama pull'"],
], [5.4, 4.2, 7], caption="Table 6. Issue log (full detail in CHANGELOG.md)")
doc.add_heading("Known failure cases, useful for the report", 3)
bullets([
    "'Is SQL knowledge expected for business analysts?' ranks first an ad that never mentions SQL but repeats 'knowledge', 'expected' and 'business analyst'.",
    "'Is the R language asked for?' matches an ad about staff speaking 70 'languages'.",
    "'Is a forklift licence required?' matches an ad about regulatory 'licence' compliance.",
    "An obvious fix, boosting ads whose skill list contains the queried skill, was rejected: our correct answers are defined from those same lists, so it would inflate the scores artificially.",
])

# ---------------- 7 Code map ----------------
doc.add_heading("7. Where everything lives", 1)
table(["File", "Lines", "What it does"], [
    ["skillgap/core.py", "231", "Loading ads, tokenising, skill list and extraction, BM25 / dense / hybrid search, skill gaps"],
    ["skillgap/evaluation.py", "153", "Metrics, threshold tuning, confidence intervals, skill-extraction scoring"],
    ["skillgap/generation.py", "57", "Prompts v1 and v2, Ollama call, citation and refusal checks"],
    ["app.py", "83", "The Streamlit web app"],
    ["ingest.py", "95", "Cleaning, deduplication, skill tagging, provenance"],
    ["scripts/build_seek_dataset.py", "81", "Download, checksum and selection of the SEEK ads"],
    ["scripts/build_questions.py", "116", "The 54 labelled questions and the two checking sheets"],
    ["run_experiments.py", "152", "Tuning, comparison table, LaTeX tables, figure, run log"],
    ["run_generation_eval.py", "73", "Runs the AI on every question for human marking, then scores it"],
    ["run_eval.py", "18", "Scores a single setting"],
    ["tests/test_core.py", "103", "15 automated tests"],
], [5.2, 1.4, 10], caption="Table 7. Code map", font=9)
para("Commands, from the skillgap_build folder after 'source .venv/bin/activate':", bold_lead=None)
for cmd in ["pytest -q                           # 15 tests", "python scripts/build_seek_dataset.py  # rebuild the 159 ads", "python scripts/build_questions.py     # rebuild the questions",
            "python run_experiments.py            # all results, about 30 seconds", "streamlit run app.py                 # the app (needs Ollama for answers)"]:
    p = doc.add_paragraph(); r = p.add_run(cmd); r.font.name = "Menlo"; r.font.size = Pt(9); p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.left_indent = Cm(0.6)
doc.add_paragraph()

# ---------------- 8 Next ----------------
doc.add_heading("8. Open items to discuss as a team", 1)
bullets([
    ("Check our labels (about 1 hour, split between two people): ", "mark the 72 rows of label_audit.csv and hand-label the 20 ads in skill_labels.csv, then re-run run_experiments.py to get skill-extraction precision and recall."),
    ("Answer-quality test: ", "run run_generation_eval.py, two people mark each answer as faithful (every claim supported) and correct, compare prompt v1 against v2 on dev, then run test once."),
    ("Limitations to state: ", "ads from one week in April 2018 (e.g. Power BI under-represented, no generative-AI skills); most ads are not graduate-specific; only 21 Data Scientist ads; small test set with wide confidence intervals; skill tagging only finds skills on our list of 47."),
    ("Business value, framed carefully: ", "the evidence shows cited, checkable answers and honest refusals on tested questions; it does not show better employment outcomes."),
    ("Project management evidence: ", "start a GitHub repository and commit regularly; log the work on Trello; record who did what for the contribution sheet."),
    ("AI declaration: ", "list the AI-assisted parts (code changes, data curation, question writing, this document) in the Condition 3 form."),
])

doc.save(OUT)
print("saved", OUT)
